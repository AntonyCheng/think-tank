from __future__ import annotations

import asyncio
import io

import fitz
import pytest
from docx import Document

from app.source_access import (
    AioHttpSourceTransport,
    SourceAccessError,
    SourceAccessLimits,
    SourceHttpResponse,
    SourceMaterializer,
    _source_http_proxy,
    source_fallback_decision,
    url_allowed_by_domains,
)


class StubResolver:
    def __init__(self, *addresses: str) -> None:
        self.addresses = addresses

    async def resolve(self, host: str, port: int) -> tuple[str, ...]:
        return self.addresses


class MappingResolver:
    def __init__(self, addresses: dict[str, tuple[str, ...]]) -> None:
        self.addresses = addresses

    async def resolve(self, host: str, port: int) -> tuple[str, ...]:
        del port
        return self.addresses[host]


class UnexpectedResolver:
    async def resolve(self, host: str, port: int) -> tuple[str, ...]:
        raise AssertionError("forbidden hostnames must not reach DNS")


class UnexpectedTransport:
    async def fetch(self, *args, **kwargs):
        raise AssertionError("unsafe addresses must not reach HTTP transport")


class StubTransport:
    def __init__(self, response: SourceHttpResponse) -> None:
        self.response = response

    async def fetch(self, *args, **kwargs) -> SourceHttpResponse:
        return self.response


class SequenceTransport:
    def __init__(self, *responses: SourceHttpResponse) -> None:
        self.responses = list(responses)
        self.calls = 0

    async def fetch(self, *args, **kwargs) -> SourceHttpResponse:
        response = self.responses[self.calls]
        self.calls += 1
        return response


def test_short_static_html_is_eligible_for_browser_recovery() -> None:
    decision = source_fallback_decision(
        b"<html><body><p>Short public notice.</p></body></html>",
        "text/html",
        "Short public notice.",
    )

    assert decision.reason == "short_text"


def test_empty_client_rendered_html_is_eligible_for_one_fallback() -> None:
    decision = source_fallback_decision(
        b"<html><body><div id=\"root\"></div><script src=\"/app.js\"></script></body></html>",
        "text/html",
        "",
    )

    assert decision.reason == "empty_text"


def test_rendered_fallback_replaces_an_empty_static_html_response() -> None:
    static = SourceHttpResponse(
        status=200,
        headers={"content-type": "text/html"},
        body=b"<html><body><div id=\"root\"></div><script src=\"/app.js\"></script></body></html>",
        peer_ip="93.184.216.34",
    )
    rendered = SourceHttpResponse(
        status=200,
        headers={"content-type": "text/html"},
        body=b"<html><head><title>Rendered</title></head><body><main>Verified rendered evidence.</main></body></html>",
        peer_ip="93.184.216.34",
    )
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(static),
        fallback_transport=StubTransport(rendered),
        fallback_strategy="browser",
    )

    result = asyncio.run(
        materializer.materialize(["https://public.example/report"])
    )

    assert result.sources[0].fetch_strategy == "browser"
    assert result.sources[0].fallback_reason == "empty_text"
    assert result.sources[0].text == "Verified rendered evidence."


def test_rendered_fallback_recovers_a_static_forbidden_response() -> None:
    static = SourceHttpResponse(
        status=403,
        headers={"content-type": "text/html"},
        body=b"Access denied",
        peer_ip="93.184.216.34",
    )
    rendered = SourceHttpResponse(
        status=200,
        headers={"content-type": "text/html"},
        body=b"<html><head><title>Recovered</title></head><body><main>Verified evidence from a browser render.</main></body></html>",
        peer_ip="93.184.216.34",
    )
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(static),
        fallback_transport=StubTransport(rendered),
        fallback_strategy="browser",
    )

    result = asyncio.run(
        materializer.materialize(["https://public.example/report"])
    )

    assert result.sources[0].fetch_strategy == "browser"
    assert result.sources[0].fallback_reason == "unsuccessful_response"
    assert "Verified evidence" in result.sources[0].text


def test_rendered_fallback_rejects_an_unverified_peer() -> None:
    static = SourceHttpResponse(
        status=200,
        headers={"content-type": "text/html"},
        body=b"<div id=\"root\"></div><script src=\"/app.js\"></script>",
        peer_ip="93.184.216.34",
    )
    rendered = SourceHttpResponse(
        status=200,
        headers={"content-type": "text/html"},
        body=b"<main>Unexpected rendered content.</main>",
        peer_ip="93.184.216.35",
    )
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(static),
        fallback_transport=StubTransport(rendered),
    )

    result = asyncio.run(
        materializer.materialize(["https://public.example/report"])
    )

    assert result.sources == ()
    assert result.failures[0].code == "source_content_empty"


def test_private_dns_result_is_rejected_before_http_access() -> None:
    materializer = SourceMaterializer(
        resolver=StubResolver("127.0.0.1"),
        transport=UnexpectedTransport(),
    )

    with pytest.raises(SourceAccessError) as captured:
        asyncio.run(materializer.materialize(["https://public.example/report"]))

    assert captured.value.code == "source_address_forbidden"
    assert captured.value.url == "https://public.example/report"


def test_public_html_is_materialized_as_canonical_text_evidence() -> None:
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(
            SourceHttpResponse(
                status=200,
                headers={"content-type": "text/html; charset=utf-8"},
                body=(
                    b"<html><head><title>Annual Report</title></head>"
                    b"<body><main>Revenue grew by 12 percent.</main>"
                    b"<script>secret()</script></body></html>"
                ),
                peer_ip="93.184.216.34",
            )
        ),
    )

    result = asyncio.run(
        materializer.materialize(["HTTPS://Public.Example:443/report"])
    )

    assert result.total_bytes == 132
    assert len(result.sources) == 1
    assert result.sources[0].canonical_url == "https://public.example/report"
    assert result.sources[0].title == "Annual Report"
    assert result.sources[0].text == "Revenue grew by 12 percent."


def test_plain_text_respects_declared_charset_without_text_truncation() -> None:
    body = "季度报告：收入增长百分之十二。".encode("gb18030")
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(
            SourceHttpResponse(
                status=200,
                headers={"content-type": "text/plain; charset=gb18030"},
                body=body,
                peer_ip="93.184.216.34",
            )
        ),
    )

    result = asyncio.run(
        materializer.materialize(["https://public.example/report.txt"])
    )

    assert result.sources[0].media_type == "text/plain"
    assert result.sources[0].title == "report.txt"
    assert result.sources[0].text == "季度报告：收入增长百分之十二。"


def test_pdf_text_is_extracted_without_a_second_network_request() -> None:
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Verified annual revenue: 42 million.")
    body = document.tobytes()
    document.close()
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(
            SourceHttpResponse(
                status=200,
                headers={"content-type": "application/pdf"},
                body=body,
                peer_ip="93.184.216.34",
            )
        ),
    )

    result = asyncio.run(
        materializer.materialize(["https://public.example/annual.pdf"])
    )

    assert result.sources[0].media_type == "application/pdf"
    assert result.sources[0].title == "annual.pdf"
    assert "Verified annual revenue: 42 million." in result.sources[0].text


def test_url_credentials_are_rejected_before_dns_resolution() -> None:
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=UnexpectedTransport(),
    )

    with pytest.raises(SourceAccessError) as captured:
        asyncio.run(
            materializer.materialize(
                ["https://user:password@public.example/report"]
            )
        )

    assert captured.value.code == "source_url_invalid"


def test_redirect_to_private_address_is_rejected_before_second_request() -> None:
    transport = SequenceTransport(
        SourceHttpResponse(
            status=302,
            headers={"location": "http://internal.example/admin"},
            body=b"",
            peer_ip="93.184.216.34",
        )
    )
    materializer = SourceMaterializer(
        resolver=MappingResolver(
            {
                "public.example": ("93.184.216.34",),
                "internal.example": ("127.0.0.1",),
            }
        ),
        transport=transport,
    )

    with pytest.raises(SourceAccessError) as captured:
        asyncio.run(
            materializer.materialize(["https://public.example/report"])
        )

    assert captured.value.code == "source_redirect_forbidden"
    assert transport.calls == 1


def test_redirect_to_a_disallowed_url_form_is_a_redirect_failure() -> None:
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(
            SourceHttpResponse(
                status=302,
                headers={"location": "file:///server/private"},
                body=b"",
                peer_ip="93.184.216.34",
            )
        ),
    )

    with pytest.raises(SourceAccessError) as captured:
        asyncio.run(
            materializer.materialize(["https://public.example/report"])
        )

    assert captured.value.code == "source_redirect_forbidden"


def test_response_body_cannot_exceed_per_source_limit() -> None:
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(
            SourceHttpResponse(
                status=200,
                headers={"content-type": "text/html"},
                body=b"x" * (5 * 1024 * 1024 + 1),
                peer_ip="93.184.216.34",
            )
        ),
    )

    with pytest.raises(SourceAccessError) as captured:
        asyncio.run(
            materializer.materialize(["https://public.example/large"])
        )

    assert captured.value.code == "source_response_too_large"


def test_materialized_sources_share_one_total_download_budget() -> None:
    response = SourceHttpResponse(
        status=200,
        headers={"content-type": "text/html"},
        body=b"<p>1234</p>",
        peer_ip="93.184.216.34",
    )
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=SequenceTransport(response, response),
        limits=SourceAccessLimits(
            max_response_bytes=20,
            max_total_bytes=20,
        ),
    )

    with pytest.raises(SourceAccessError) as captured:
        asyncio.run(
            materializer.materialize(
                [
                    "https://public.example/one",
                    "https://public.example/two",
                ]
            )
        )

    assert captured.value.code == "source_total_too_large"


def test_equivalent_urls_are_fetched_once_after_normalization() -> None:
    transport = SequenceTransport(
        SourceHttpResponse(
            status=200,
            headers={"content-type": "text/html"},
            body=b"<p>Evidence</p>",
            peer_ip="93.184.216.34",
        )
    )
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=transport,
    )

    result = asyncio.run(
        materializer.materialize(
            [
                "HTTPS://Public.Example:443/report",
                "https://public.example/report",
            ]
        )
    )

    assert len(result.sources) == 1
    assert transport.calls == 1


def test_domain_allowlist_does_not_accept_a_lookalike_suffix() -> None:
    assert url_allowed_by_domains(
        "https://news.example.com/article",
        include_domains=("example.com",),
    )
    assert not url_allowed_by_domains(
        "https://example.com.evil.test/article",
        include_domains=("example.com",),
    )


def test_localhost_is_rejected_before_dns_resolution() -> None:
    materializer = SourceMaterializer(
        resolver=UnexpectedResolver(),
        transport=UnexpectedTransport(),
    )

    with pytest.raises(SourceAccessError) as captured:
        asyncio.run(materializer.materialize(["http://localhost/admin"]))

    assert captured.value.code == "source_address_forbidden"


def test_dns_transport_errors_use_a_stable_public_failure() -> None:
    class FailingResolver:
        async def resolve(self, host: str, port: int) -> tuple[str, ...]:
            del host, port
            raise OSError("private resolver detail")

    materializer = SourceMaterializer(
        resolver=FailingResolver(),
        transport=UnexpectedTransport(),
    )

    result = asyncio.run(
        materializer.materialize(["https://public.example/report"])
    )

    assert result.sources == ()
    assert result.failures[0].code == "source_unavailable"
    assert "private resolver detail" not in result.failures[0].message


def test_ordinary_unavailable_source_is_recorded_while_others_continue() -> None:
    resolver = MappingResolver({
        "missing.example": ("93.184.216.34",),
        "public.example": ("93.184.216.35",),
    })
    transport = SequenceTransport(
        SourceHttpResponse(
            status=503,
            headers={"content-type": "text/plain"},
            body=b"unavailable",
            peer_ip="93.184.216.34",
        ),
        SourceHttpResponse(
            status=200,
            headers={"content-type": "text/plain"},
            body=b"verified evidence",
            peer_ip="93.184.216.35",
        ),
    )
    materializer = SourceMaterializer(
        resolver=resolver,
        transport=transport,
    )

    result = asyncio.run(
        materializer.materialize([
            "https://missing.example/report",
            "https://public.example/report",
        ])
    )

    assert [source.canonical_url for source in result.sources] == [
        "https://public.example/report"
    ]
    assert [failure.code for failure in result.failures] == [
        "source_unavailable"
    ]
    assert result.failures[0].url == "https://missing.example/report"


def test_office_document_url_is_converted_to_markdown_without_a_hint() -> None:
    document = Document()
    document.add_heading("Quarterly Update", level=1)
    document.add_paragraph("Revenue rose to 42 million euros in Q3.")
    buffer = io.BytesIO()
    document.save(buffer)
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(
            SourceHttpResponse(
                status=200,
                headers={
                    "content-type": (
                        "application/vnd.openxmlformats-officedocument"
                        ".wordprocessingml.document"
                    )
                },
                body=buffer.getvalue(),
                peer_ip="93.184.216.34",
            )
        ),
    )

    result = asyncio.run(
        materializer.materialize(["https://public.example/q3.docx"])
    )

    assert "Quarterly Update" in result.sources[0].text
    assert "42 million euros" in result.sources[0].text


def test_csv_url_is_rendered_as_a_markdown_table() -> None:
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(
            SourceHttpResponse(
                status=200,
                headers={"content-type": "text/csv"},
                body=b"name,amount\nAlpha,10\nBeta,20\n",
                peer_ip="93.184.216.34",
            )
        ),
    )

    result = asyncio.run(
        materializer.materialize(["https://public.example/ledger.csv"])
    )

    assert "| name | amount |" in result.sources[0].text
    assert "| Alpha | 10 |" in result.sources[0].text


def test_html_boilerplate_is_dropped_in_favor_of_the_article_body() -> None:
    article_text = (
        "The union secured a four percent wage increase for warehouse staff "
        "after three rounds of bargaining. " * 4
    )
    body = (
        "<html><head><title>Wage settlement</title></head><body>"
        "<nav><a href='/'>Home</a><a href='/news'>News</a></nav>"
        "<header>Site masthead and search</header>"
        f"<article><p>{article_text}</p></article>"
        "<aside class='related'>Related coverage you might like</aside>"
        "<footer>Copyright and cookie policy</footer>"
        "</body></html>"
    ).encode("utf-8")
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(
            SourceHttpResponse(
                status=200,
                headers={"content-type": "text/html; charset=utf-8"},
                body=body,
                peer_ip="93.184.216.34",
            )
        ),
    )

    result = asyncio.run(
        materializer.materialize(["https://public.example/story"])
    )

    text = result.sources[0].text
    assert "four percent wage increase" in text
    assert "Related coverage" not in text
    assert "cookie policy" not in text
    assert "Home" not in text


def test_nested_navigation_regions_do_not_break_extraction() -> None:
    body = (
        "<html><head><title>Filing</title></head><body>"
        "<div role='navigation'><nav role='search'><a href='/'>Home</a></nav>"
        "<ul role='menu'><li>More</li></ul></div>"
        "<main><p>Revenue rose twelve percent to forty two million euros in the "
        "quarter, according to the audited filing released today.</p></main>"
        "<div class='footer social-share'>Share this article</div>"
        "</body></html>"
    ).encode("utf-8")
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(
            SourceHttpResponse(
                status=200,
                headers={"content-type": "text/html; charset=utf-8"},
                body=body,
                peer_ip="93.184.216.34",
            )
        ),
    )

    result = asyncio.run(
        materializer.materialize(["https://public.example/filing"])
    )

    text = result.sources[0].text
    assert "forty two million euros" in text
    assert "Home" not in text
    assert "Share this article" not in text


def test_large_markup_with_little_text_is_eligible_for_a_render_fallback() -> None:
    visible_text = "This paragraph is the only readable sentence on the page. " * 6
    decision = source_fallback_decision(
        b"<div>" + b"<span></span>" * 12_000 + b"</div>",
        "text/html",
        visible_text,
    )

    assert len(visible_text.strip()) > 200
    assert decision.reason == "low_text_ratio"


def test_client_rendered_shell_is_eligible_for_a_render_fallback() -> None:
    shell_text = (
        "Loading the application, please wait a moment while the dashboard "
        "and its saved views are prepared. " * 3
    )
    decision = source_fallback_decision(
        b"<html><body><div id=\"root\"></div>"
        b"<script>window.__NEXT_DATA__={}</script></body></html>",
        "text/html",
        shell_text,
    )

    assert 200 < len(shell_text.strip()) < 1_200
    assert decision.reason == "client_rendered_shell"


def test_source_http_proxy_env_is_validated(monkeypatch) -> None:
    monkeypatch.delenv("GPTR_SOURCE_HTTP_PROXY", raising=False)
    monkeypatch.delenv("GPTR_SOURCE_HTTPS_PROXY", raising=False)
    assert _source_http_proxy() is None

    monkeypatch.setenv("GPTR_SOURCE_HTTP_PROXY", "not-a-url")
    assert _source_http_proxy() is None

    monkeypatch.setenv("GPTR_SOURCE_HTTP_PROXY", "http://egress.proxy.internal:3128")
    assert _source_http_proxy() == "http://egress.proxy.internal:3128"


def test_proxy_fetch_skips_the_socket_peer_pin() -> None:
    class ProxyTransport:
        async def fetch(self, *args, **kwargs) -> SourceHttpResponse:
            # The observed peer is the proxy, never one of the pinned addresses.
            return SourceHttpResponse(
                status=200,
                headers={"content-type": "text/plain"},
                body=b"verified evidence via proxy",
                peer_ip="10.4.4.4",
                via_proxy=True,
            )

    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=ProxyTransport(),
    )

    result = asyncio.run(
        materializer.materialize(["https://public.example/report"])
    )

    assert result.sources[0].text == "verified evidence via proxy"


def test_transport_without_proxy_still_enforces_the_peer_pin() -> None:
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(
            SourceHttpResponse(
                status=200,
                headers={"content-type": "text/plain"},
                body=b"evidence from an unexpected address",
                peer_ip="203.0.113.9",
            )
        ),
    )

    with pytest.raises(SourceAccessError) as captured:
        asyncio.run(
            materializer.materialize(["https://public.example/report"])
        )

    assert captured.value.code == "source_peer_forbidden"


def test_aiohttp_transport_accepts_an_explicit_proxy() -> None:
    assert AioHttpSourceTransport(proxy="http://p:3128")._proxy == "http://p:3128"
    assert AioHttpSourceTransport()._proxy is None


def test_html_mislabeled_as_octet_stream_is_still_read_as_a_page() -> None:
    materializer = SourceMaterializer(
        resolver=StubResolver("93.184.216.34"),
        transport=StubTransport(
            SourceHttpResponse(
                status=200,
                headers={"content-type": "application/octet-stream"},
                body=(
                    b"<!DOCTYPE html><html><head><title>Notice</title></head>"
                    b"<body><main>The board approved the merger on Tuesday.</main>"
                    b"</body></html>"
                ),
                peer_ip="93.184.216.34",
            )
        ),
    )

    result = asyncio.run(
        materializer.materialize(["https://public.example/notice"])
    )

    assert result.sources[0].title == "Notice"
    assert result.sources[0].text == "The board approved the merger on Tuesday."
