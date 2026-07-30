from zipfile import ZipFile

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from app.exporter import export_docx, export_markdown, export_pdf
from app.main import app


SAMPLE = """# 城市人工智能产业研究

## 核心结论

产业规模达到 100 亿元[\\[1\\]](https://example.com/data)。

| 指标 | 数值 |
| --- | --- |
| 企业数量 | 42 家 |

## 参考来源

[1] [公开统计数据](https://example.com/data)
"""


def test_exports_markdown_docx_and_pdf(tmp_path) -> None:
    markdown_path = export_markdown(SAMPLE, "task-1", tmp_path)
    docx_path = export_docx("测试研究", SAMPLE, "task-1", tmp_path)
    pdf_path = export_pdf("测试研究", SAMPLE, "task-1", tmp_path)

    assert markdown_path.read_text(encoding="utf-8") == SAMPLE
    assert docx_path.read_bytes().startswith(b"PK")
    assert pdf_path.read_bytes().startswith(b"%PDF")
    pdf_text = "\n".join(
        page.extract_text() or "" for page in PdfReader(pdf_path).pages
    )
    assert "城市人工智能产业研究" in pdf_text
    assert "产业规模达到 100 亿元" in pdf_text
    assert "[1]" in pdf_text
    assert "THINK TANK" not in pdf_text

    document = Document(docx_path)
    assert document.core_properties.author != "Think Tank"
    text = "\n".join(
        paragraph.text for paragraph in document.paragraphs
    )
    assert "城市人工智能产业研究" in text
    assert "核心结论" in text
    assert "[1]" in text
    assert "THINK TANK" not in text
    header_text = "\n".join(
        paragraph.text
        for section in document.sections
        for paragraph in section.header.paragraphs
    )
    assert "THINK TANK" not in header_text
    assert len(document.tables) == 1
    with ZipFile(docx_path) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "w:hyperlink" in xml
    assert "superscript" in xml


def test_export_endpoint_returns_download() -> None:
    response = TestClient(app).post(
        "/export/docx",
        json={
            "taskId": "endpoint-test",
            "title": "测试研究",
            "markdown": SAMPLE,
        },
    )

    assert response.status_code == 200
    assert response.content.startswith(b"PK")
    assert "think-tank-report.docx" in response.headers[
        "content-disposition"
    ]
