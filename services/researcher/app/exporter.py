from __future__ import annotations

import html
import re
from datetime import date
from pathlib import Path
from typing import Any

import mistune
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.fonts import addMapping
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    ListFlowable,
    ListItem,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


INK = RGBColor(0x0B, 0x25, 0x45)
HEADING = RGBColor(0x2E, 0x74, 0xB5)
HEADING_DARK = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x66, 0x70, 0x68)
BODY_FONT = "Microsoft YaHei"
FALLBACK_FONT = "Arial"
PDF_FONT = "ThinkTankNotoSansSCMedium"
PDF_FONT_PATH = (
    Path(__file__).resolve().parent.parent
    / "assets"
    / "fonts"
    / "NotoSansSC-Medium.ttf"
)
PDF_TEXT_COLOR = "#000000"
MARKDOWN = mistune.create_markdown(renderer="ast", plugins=["table"])


def export_markdown(
    markdown: str,
    task_id: str,
    export_root: Path,
) -> Path:
    path = _export_path(task_id, export_root, "md")
    path.write_text(markdown, encoding="utf-8")
    return path


def export_docx(
    title: str,
    markdown: str,
    task_id: str,
    export_root: Path,
) -> Path:
    tokens = MARKDOWN(markdown)
    document_title, body_tokens = _extract_title(title, tokens)
    path = _export_path(task_id, export_root, "docx")

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    _configure_docx_styles(doc)
    _add_page_footer(section)
    _add_docx_title(doc, document_title, title)

    for token in body_tokens:
        _append_docx_block(doc, token)

    doc.core_properties.title = document_title
    doc.core_properties.subject = title
    doc.core_properties.author = ""
    doc.core_properties.keywords = "research, verified sources"
    doc.save(path)
    return path


def export_pdf(
    title: str,
    markdown: str,
    task_id: str,
    export_root: Path,
) -> Path:
    tokens = MARKDOWN(markdown)
    document_title, body_tokens = _extract_title(title, tokens)
    path = _export_path(task_id, export_root, "pdf")
    styles = _pdf_styles()
    story: list[Any] = [
        Spacer(1, 0.4 * inch),
        Paragraph(html.escape(document_title), styles["Title"]),
        Paragraph(
            f"研究主题：{html.escape(title)}<br/>生成日期：{date.today().isoformat()}",
            styles["Metadata"],
        ),
        Spacer(1, 0.22 * inch),
    ]
    for token in body_tokens:
        story.extend(_pdf_block(token, styles))

    pdf = SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        rightMargin=inch,
        leftMargin=inch,
        topMargin=0.85 * inch,
        bottomMargin=0.75 * inch,
        title=document_title,
        author="",
    )
    pdf.build(
        story,
        onFirstPage=_draw_pdf_page,
        onLaterPages=_draw_pdf_page,
    )
    return path


def _export_path(task_id: str, export_root: Path, extension: str) -> Path:
    safe_id = re.sub(r"[^a-zA-Z0-9_-]", "-", task_id)[:80] or "report"
    directory = export_root / safe_id
    directory.mkdir(parents=True, exist_ok=True)
    return directory / f"report.{extension}"


def _extract_title(
    fallback: str,
    tokens: list[dict[str, Any]],
) -> tuple[str, list[dict[str, Any]]]:
    body = list(tokens)
    for index, token in enumerate(body):
        if token.get("type") == "heading" and token.get("attrs", {}).get("level") == 1:
            title = _plain_text(token.get("children", []))
            del body[index]
            return title or fallback, body
    return fallback, body


def _configure_docx_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    _set_style_font(normal, BODY_FONT, 11, RGBColor(0, 0, 0))
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    heading_tokens = {
        "Heading 1": (16, HEADING, 16, 8),
        "Heading 2": (13, HEADING, 12, 6),
        "Heading 3": (12, HEADING_DARK, 8, 4),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[name]
        _set_style_font(style, BODY_FONT, size, color, bold=True)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        _set_style_font(style, BODY_FONT, 11, RGBColor(0, 0, 0))
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def _set_style_font(
    style: Any,
    name: str,
    size: float,
    color: RGBColor,
    bold: bool = False,
) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold
    style._element.rPr.rFonts.set(qn("w:ascii"), FALLBACK_FONT)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), FALLBACK_FONT)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _add_page_footer(section: Any) -> None:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("第 ")
    _set_run_font(run, 9, MUTED)
    _append_field(run, "PAGE")
    run = footer.add_run(" 页")
    _set_run_font(run, 9, MUTED)


def _add_docx_title(doc: Document, title: str, topic: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(title)
    _set_run_font(run, 24, INK, bold=True)

    metadata = doc.add_paragraph()
    metadata.paragraph_format.space_after = Pt(16)
    metadata.paragraph_format.line_spacing = 1.15
    run = metadata.add_run(
        f"研究主题：{topic}\n生成日期：{date.today().isoformat()}"
    )
    _set_run_font(run, 9.5, MUTED)


def _append_docx_block(doc: Document, token: dict[str, Any]) -> None:
    kind = token.get("type")
    if kind == "blank_line":
        return
    if kind == "heading":
        level = min(max(int(token.get("attrs", {}).get("level", 2)) - 1, 1), 3)
        paragraph = doc.add_paragraph(style=f"Heading {level}")
        _append_docx_inline(paragraph, token.get("children", []))
        return
    if kind in {"paragraph", "block_text"}:
        paragraph = doc.add_paragraph()
        _append_docx_inline(paragraph, token.get("children", []))
        return
    if kind == "block_quote":
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.space_before = Pt(4)
        paragraph.paragraph_format.space_after = Pt(8)
        _paragraph_left_border(paragraph, "2E74B5")
        for child in token.get("children", []):
            _append_docx_inline(paragraph, child.get("children", []))
        return
    if kind == "list":
        style = "List Number" if token.get("attrs", {}).get("ordered") else "List Bullet"
        for item in token.get("children", []):
            paragraph = doc.add_paragraph(style=style)
            for child in item.get("children", []):
                _append_docx_inline(paragraph, child.get("children", []))
        return
    if kind == "table":
        _append_docx_table(doc, token)
        return
    if kind == "thematic_break":
        paragraph = doc.add_paragraph()
        _paragraph_bottom_border(paragraph, "D7DBE2")
        return
    if kind == "block_code":
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.space_after = Pt(8)
        run = paragraph.add_run(token.get("raw", ""))
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        return


def _append_docx_inline(paragraph: Any, tokens: list[dict[str, Any]]) -> None:
    for token in tokens:
        kind = token.get("type")
        if kind == "text":
            run = paragraph.add_run(token.get("raw", ""))
            _set_run_font(run, 11, RGBColor(0, 0, 0))
        elif kind in {"strong", "emphasis", "codespan", "strikethrough"}:
            text = _plain_text(token.get("children", [])) or token.get("raw", "")
            run = paragraph.add_run(text)
            _set_run_font(run, 10 if kind == "codespan" else 11, RGBColor(0, 0, 0))
            run.bold = kind == "strong"
            run.italic = kind == "emphasis"
            run.font.strike = kind == "strikethrough"
            if kind == "codespan":
                run.font.name = "Consolas"
        elif kind == "link":
            label = _plain_text(token.get("children", []))
            url = str(token.get("attrs", {}).get("url", ""))
            citation = re.fullmatch(r"\[(\d+)\]", label)
            _add_hyperlink(
                paragraph,
                label,
                url,
                superscript=bool(citation),
            )
        elif kind in {"softbreak", "linebreak"}:
            paragraph.add_run().add_break()
        elif token.get("children"):
            _append_docx_inline(paragraph, token["children"])


def _append_docx_table(doc: Document, token: dict[str, Any]) -> None:
    rows: list[list[dict[str, Any]]] = []
    for section in token.get("children", []):
        if section.get("type") == "table_head":
            rows.append(section.get("children", []))
        elif section.get("type") == "table_body":
            rows.extend(
                row.get("children", [])
                for row in section.get("children", [])
            )
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Table Grid"
    widths = [9360 // columns] * columns
    widths[-1] += 9360 - sum(widths)
    _set_table_geometry(table, widths)
    for row_index, row in enumerate(rows):
        for column_index, cell_token in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            _append_docx_inline(paragraph, cell_token.get("children", []))
            if row_index == 0:
                _shade_cell(cell, "F2F4F7")
                for run in paragraph.runs:
                    run.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _set_table_geometry(table: Any, widths: list[int]) -> None:
    properties = table._tbl.tblPr
    width = properties.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        properties.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), "9360")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "120")
    properties.append(indent)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for column_width in widths:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(column_width))
        grid.append(grid_column)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index] / 1440)
            cell_properties = cell._tc.get_or_add_tcPr()
            cell_width = cell_properties.first_child_found_in("w:tcW")
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                cell_properties.append(cell_width)
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(widths[index]))


def _add_hyperlink(
    paragraph: Any,
    text: str,
    url: str,
    *,
    superscript: bool = False,
) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E74B5")
    properties.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.append(underline)
    if superscript:
        vertical = OxmlElement("w:vertAlign")
        vertical.set(qn("w:val"), "superscript")
        properties.append(vertical)
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _append_field(run: Any, field: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = field
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _set_run_font(
    run: Any,
    size: float,
    color: RGBColor,
    *,
    bold: bool = False,
) -> None:
    run.font.name = BODY_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FALLBACK_FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FALLBACK_FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold


def _paragraph_left_border(paragraph: Any, color: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        properties.append(borders)
    border = OxmlElement("w:left")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "18")
    border.set(qn("w:space"), "8")
    border.set(qn("w:color"), color)
    borders.append(border)


def _paragraph_bottom_border(paragraph: Any, color: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    border = OxmlElement("w:bottom")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), "6")
    border.set(qn("w:color"), color)
    borders.append(border)
    properties.append(borders)


def _shade_cell(cell: Any, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _pdf_styles() -> dict[str, ParagraphStyle]:
    _register_pdf_fonts()
    styles = getSampleStyleSheet()
    base = {
        "fontName": PDF_FONT,
    }
    return {
        "Title": ParagraphStyle(
            "ResearchTitle",
            parent=styles["Title"],
            fontName=PDF_FONT,
            fontSize=24,
            leading=31,
            textColor=colors.HexColor(PDF_TEXT_COLOR),
            alignment=TA_LEFT,
            spaceAfter=8,
        ),
        "Metadata": ParagraphStyle(
            "Metadata",
            **base,
            fontSize=9,
            leading=14,
            textColor=colors.HexColor(PDF_TEXT_COLOR),
            spaceAfter=12,
        ),
        "Body": ParagraphStyle(
            "Body",
            **base,
            fontSize=11,
            leading=16,
            textColor=colors.HexColor(PDF_TEXT_COLOR),
            spaceAfter=8,
        ),
        "H1": ParagraphStyle(
            "H1",
            **base,
            fontSize=16,
            leading=21,
            textColor=colors.HexColor(PDF_TEXT_COLOR),
            spaceBefore=14,
            spaceAfter=8,
            keepWithNext=True,
        ),
        "H2": ParagraphStyle(
            "H2",
            **base,
            fontSize=13,
            leading=18,
            textColor=colors.HexColor(PDF_TEXT_COLOR),
            spaceBefore=11,
            spaceAfter=6,
            keepWithNext=True,
        ),
        "H3": ParagraphStyle(
            "H3",
            **base,
            fontSize=11.5,
            leading=16,
            textColor=colors.HexColor(PDF_TEXT_COLOR),
            spaceBefore=8,
            spaceAfter=4,
            keepWithNext=True,
        ),
        "Quote": ParagraphStyle(
            "Quote",
            **base,
            fontSize=10,
            leading=15,
            textColor=colors.HexColor(PDF_TEXT_COLOR),
            leftIndent=18,
            borderColor=colors.HexColor(PDF_TEXT_COLOR),
            borderWidth=1.5,
            borderPadding=7,
            backColor=colors.HexColor("#F4F6F9"),
            spaceAfter=8,
        ),
        "Code": ParagraphStyle(
            "Code",
            fontName="Courier-Bold",
            fontSize=8.5,
            leading=12,
            textColor=colors.HexColor(PDF_TEXT_COLOR),
            leftIndent=12,
            backColor=colors.HexColor("#F2F4F7"),
            borderPadding=6,
            spaceAfter=8,
        ),
    }


def _pdf_block(
    token: dict[str, Any],
    styles: dict[str, ParagraphStyle],
) -> list[Any]:
    kind = token.get("type")
    if kind == "blank_line":
        return []
    if kind == "heading":
        level = min(max(int(token.get("attrs", {}).get("level", 2)) - 1, 1), 3)
        return [Paragraph(_pdf_inline(token.get("children", [])), styles[f"H{level}"])]
    if kind in {"paragraph", "block_text"}:
        return [Paragraph(_pdf_inline(token.get("children", [])), styles["Body"])]
    if kind == "block_quote":
        text = " ".join(
            _pdf_inline(child.get("children", []))
            for child in token.get("children", [])
        )
        return [Paragraph(text, styles["Quote"])]
    if kind == "list":
        items = []
        for item in token.get("children", []):
            text = " ".join(
                _pdf_inline(child.get("children", []))
                for child in item.get("children", [])
            )
            items.append(ListItem(Paragraph(text, styles["Body"])))
        return [
            ListFlowable(
                items,
                bulletType="1" if token.get("attrs", {}).get("ordered") else "bullet",
                leftIndent=28,
                bulletFontName=PDF_FONT,
                bulletFontSize=9,
                spaceAfter=8,
            )
        ]
    if kind == "table":
        return [_pdf_table(token, styles)]
    if kind == "thematic_break":
        return [Spacer(1, 8)]
    if kind == "block_code":
        return [Paragraph(html.escape(token.get("raw", "")), styles["Code"])]
    return []


def _pdf_table(
    token: dict[str, Any],
    styles: dict[str, ParagraphStyle],
) -> Table:
    rows: list[list[Any]] = []
    for section in token.get("children", []):
        if section.get("type") == "table_head":
            rows.append([
                Paragraph(_pdf_inline(cell.get("children", [])), styles["Body"])
                for cell in section.get("children", [])
            ])
        elif section.get("type") == "table_body":
            for row in section.get("children", []):
                rows.append([
                    Paragraph(_pdf_inline(cell.get("children", [])), styles["Body"])
                    for cell in row.get("children", [])
                ])
    columns = max((len(row) for row in rows), default=1)
    table = Table(
        rows,
        colWidths=[6.5 * inch / columns] * columns,
        repeatRows=1,
        hAlign="LEFT",
    )
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F2F4F7")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#C8CDD3")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    return table


def _pdf_inline(tokens: list[dict[str, Any]]) -> str:
    parts: list[str] = []
    for token in tokens:
        kind = token.get("type")
        if kind == "text":
            parts.append(html.escape(token.get("raw", "")))
        elif kind == "strong":
            parts.append(f"<b>{_pdf_inline(token.get('children', []))}</b>")
        elif kind == "emphasis":
            parts.append(f"<i>{_pdf_inline(token.get('children', []))}</i>")
        elif kind == "codespan":
            parts.append(f"<font name=\"Courier\">{html.escape(token.get('raw', ''))}</font>")
        elif kind == "link":
            label = _plain_text(token.get("children", []))
            url = html.escape(str(token.get("attrs", {}).get("url", "")), quote=True)
            if re.fullmatch(r"\[\d+\]", label):
                parts.append(
                    f"<super><a color=\"{PDF_TEXT_COLOR}\" href=\"{url}\">"
                    f"{html.escape(label)}</a></super>"
                )
            else:
                parts.append(
                    f"<a color=\"{PDF_TEXT_COLOR}\" href=\"{url}\">{html.escape(label)}</a>"
                )
        elif kind in {"softbreak", "linebreak"}:
            parts.append("<br/>")
        elif token.get("children"):
            parts.append(_pdf_inline(token["children"]))
    return "".join(parts)


def _plain_text(tokens: list[dict[str, Any]]) -> str:
    return "".join(
        token.get("raw", "")
        if token.get("type") in {"text", "codespan"}
        else _plain_text(token.get("children", []))
        for token in tokens
    )


def _draw_pdf_page(canvas: Any, document: Any) -> None:
    canvas.saveState()
    canvas.setFont(PDF_FONT, 8)
    canvas.setFillColor(colors.HexColor(PDF_TEXT_COLOR))
    canvas.drawRightString(
        7.5 * inch,
        0.45 * inch,
        f"第 {document.page} 页",
    )
    canvas.restoreState()


def _register_pdf_fonts() -> None:
    if PDF_FONT in pdfmetrics.getRegisteredFontNames():
        return
    if not PDF_FONT_PATH.is_file():
        raise RuntimeError(
            f"bundled PDF font is missing: {PDF_FONT_PATH}"
        )
    pdfmetrics.registerFont(TTFont(PDF_FONT, PDF_FONT_PATH))
    for bold in (0, 1):
        for italic in (0, 1):
            addMapping(PDF_FONT, bold, italic, PDF_FONT)
